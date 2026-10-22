import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen slides
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional color palette (SOC / Cyber theme)
    DARK_BG = RGBColor(11, 15, 23)      # #0B0F17
    CARD_BG = RGBColor(19, 27, 42)      # #131B2A
    CARD_BORDER = RGBColor(30, 41, 59)  # #1E293B
    PRIMARY_INDIGO = RGBColor(99, 102, 241) # #6366F1
    ACCENT_CYAN = RGBColor(56, 189, 248)    # #38BDF8
    TEXT_WHITE = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ALERT_RED = RGBColor(239, 68, 68)
    SUCCESS_GREEN = RGBColor(34, 197, 94)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="MDC PROJECT EVALUATION"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = PRIMARY_INDIGO

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative accent card
    add_card(s1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1), bg_color=CARD_BG, border_color=PRIMARY_INDIGO)

    # Title box
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "MAJOR DEGREE / COURSE PROJECT PRESENTATION"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Automated Phishing Incident\nInvestigation Workbench"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "A Full-Stack SOC Incident Triage Platform with Header Forensics, Anchor Spoof Detection, Threat Intel & MITRE ATT&CK Mapping"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "Presenter: Shivangi Gupta   |   Domain: Cybersecurity / SOC Operations   |   Status: Fully Deployed & Operational"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = PRIMARY_INDIGO

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement & Industry Context")

    # Card 1: The Phishing Epidemic
    add_card(s2, Inches(0.8), Inches(1.7), Inches(3.6), Inches(5.0))
    tb1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(3.2), Inches(4.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.paragraphs[0].text = "🚨 The Threat Landscape"
    tf1.paragraphs[0].font.size = Pt(16)
    tf1.paragraphs[0].font.bold = True
    tf1.paragraphs[0].font.color.rgb = ALERT_RED
    
    p = tf1.add_paragraph()
    p.text = "\n• 90%+ of cyberattacks originate via malicious phishing emails.\n\n• Attackers leverage subtle deception: brand typosquatting, hidden double extensions, and disguised hyperlinks.\n\n• High volume of daily alerts overwhelms security analysts."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # Card 2: Current Manual Pain Points
    add_card(s2, Inches(4.8), Inches(1.7), Inches(3.6), Inches(5.0))
    tb2 = s2.shapes.add_textbox(Inches(5.0), Inches(1.9), Inches(3.2), Inches(4.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.paragraphs[0].text = "⏳ The Manual Bottleneck"
    tf2.paragraphs[0].font.size = Pt(16)
    tf2.paragraphs[0].font.bold = True
    tf2.paragraphs[0].font.color.rgb = ACCENT_CYAN
    
    p = tf2.add_paragraph()
    p.text = "\n• SOC analysts manually inspect raw RFC 822 text files for SPF/DKIM/DMARC.\n\n• Requires opening 5+ separate tabs (VirusTotal, URLhaus, AbuseIPDB, WHOIS).\n\n• Average manual triage time: 15–20 minutes per email, creating fatigue and delayed incident response."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # Card 3: Our Solution
    add_card(s2, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.0), border_color=SUCCESS_GREEN)
    tb3 = s2.shapes.add_textbox(Inches(9.0), Inches(1.9), Inches(3.3), Inches(4.6))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.paragraphs[0].text = "🎯 Proposed Solution"
    tf3.paragraphs[0].font.size = Pt(16)
    tf3.paragraphs[0].font.bold = True
    tf3.paragraphs[0].font.color.rgb = SUCCESS_GREEN
    
    p = tf3.add_paragraph()
    p.text = "\n• An automated workbench that ingests any .eml file or raw MIME text.\n\n• Deconstructs headers, scans URLs for anchor spoofing, and inspects attachment magic bytes.\n\n• Enriches indicators in parallel and delivers complete risk verdict + audit report in < 2 seconds."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 3: PROJECT OBJECTIVES & SCOPE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Project Objectives & Scope of Work")

    obj_items = [
        ("1. Automated Email Header Forensics", "Decode RFC 822 & MIME multipart structures. Reconstruct reverse-chronological relay hops across Received: headers and parse SPF, DKIM, and DMARC alignment."),
        ("2. Advanced Hyperlink & Anchor Deception Analysis", "Extract all URLs from HTML/text. Identify Anchor Text Spoofing (where display text shows trusted brand like paypal.com but links to an unauthorized external IP)."),
        ("3. Deep Attachment & Binary Inspection", "Analyze magic bytes (PE executables MZ, PDFs, ZIPs), detect double extensions (.pdf.exe), and calculate Shannon Entropy (>7.1 indicating packed/encrypted payloads)."),
        ("4. Threat Intelligence & Typosquatting Engine", "Implement offline Levenshtein distance for brand lookalike detection, IDN homographs (xn--), plus parallel integration with URLhaus, AbuseIPDB, and VirusTotal."),
        ("5. Standardized Incident Documentation", "Map findings to MITRE ATT&CK® Enterprise techniques, generate auto-timelines, manage case references (CASE-XXXXXXXX), and export standalone HTML audit reports.")
    ]

    for i, (title, desc) in enumerate(obj_items):
        y = 1.7 + (i * 1.0)
        add_card(s3, Inches(0.8), Inches(y), Inches(11.7), Inches(0.85))
        tb = s3.shapes.add_textbox(Inches(1.0), Inches(y + 0.05), Inches(11.3), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_INDIGO
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE & METHODOLOGY
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "System Architecture & 4-Stage Methodology")

    stages = [
        ("STAGE 1: INGESTION", "Raw .EML or RFC822 Input", "• Python email policy\n• Multipart MIME decode\n• Body text & HTML split\n• Attachment extraction", PRIMARY_INDIGO),
        ("STAGE 2: FORENSICS", "Deep Analytical Inspection", "• Hop tracer & Origin IP\n• SPF/DKIM/DMARC audit\n• Anchor spoof detector\n• Magic bytes & Entropy", ACCENT_CYAN),
        ("STAGE 3: ENRICHMENT", "Threat Intel & Scoring", "• Levenshtein lookalikes\n• URLhaus active malware\n• AbuseIPDB abuse scores\n• 0-100 Weighted Risk", ALERT_RED),
        ("STAGE 4: OPERATION", "Triage & Audit Reporting", "• MITRE ATT&CK Matrix\n• Incident Timeline chain\n• Case management\n• HTML Report Export", SUCCESS_GREEN)
    ]

    for i, (stg, sub, bullets, color) in enumerate(stages):
        x = 0.8 + (i * 2.95)
        add_card(s4, Inches(x), Inches(1.7), Inches(2.8), Inches(5.0))
        tb = s4.shapes.add_textbox(Inches(x + 0.15), Inches(1.85), Inches(2.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = stg
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = color
        
        p1 = tf.add_paragraph()
        p1.text = sub
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(12)
        
        p2 = tf.add_paragraph()
        p2.text = bullets
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: CORE ALGORITHMS & IMPLEMENTATION
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Key Algorithms & Code Implementation")

    # Box 1: Anchor Text Spoofing
    add_card(s5, Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.4))
    tb1 = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(2.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. Anchor Text Spoofing Detection"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p2 = tf1.add_paragraph()
    p2.text = "Compares domain in visible text against the actual destination href.\nExample: text says 'paypal.com' but href targets '194.87.139.12'.\nFormula: claimed_domain != actual_domain → Flag Spoofing"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

    # Box 2: Shannon Entropy
    add_card(s5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.4))
    tb2 = s5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(2.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. Shannon Entropy for Packed Payloads"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ALERT_RED
    p2 = tf2.add_paragraph()
    p2.text = "Calculates binary randomness density (0.0 to 8.0).\nHigh values (>7.1) indicate packed/encrypted malware bypassing antivirus.\nFormula: H(X) = - Σ P(x_i) * log2(P(x_i))"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

    # Box 3: Levenshtein Distance
    add_card(s5, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.4))
    tb3 = s5.shapes.add_textbox(Inches(1.0), Inches(4.4), Inches(5.3), Inches(2.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. Levenshtein Distance (Typosquatting)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p2 = tf3.add_paragraph()
    p2.text = "Measures edit operations between extracted domain and target brands.\nCatches lookalikes like 'update-paypal-security.cc' & punycode (xn--).\nRuns 100% offline with zero API rate-limit delays."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

    # Box 4: Transparent Risk Scorer
    add_card(s5, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.4))
    tb4 = s5.shapes.add_textbox(Inches(7.0), Inches(4.4), Inches(5.3), Inches(2.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "4. Weighted Risk Scoring Engine"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN
    p2 = tf4.add_paragraph()
    p2.text = "Point-based audit: SPF fail (+25), DKIM fail (+20), Anchor spoof (+30), Executable attachment (+35), URLhaus malware (+40).\nVerdicts: >=70 MALICIOUS | >=35 SUSPICIOUS | <35 BENIGN"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 6: LIVE PROTOTYPE & DASHBOARD UI
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Working Prototype: SOC Analyst Workbench UI", category="DEMONSTRATION & PROTOTYPE")

    # Add dashboard snapshot
    dash_img = "docs/images/dashboard_overview.png"
    if os.path.exists(dash_img):
        s6.shapes.add_picture(dash_img, Inches(0.8), Inches(1.7), Inches(7.5), Inches(5.0))

    # Notes card on the right
    add_card(s6, Inches(8.6), Inches(1.7), Inches(3.9), Inches(5.0))
    tb_proto = s6.shapes.add_textbox(Inches(8.8), Inches(1.9), Inches(3.5), Inches(4.6))
    tf_p = tb_proto.text_frame
    tf_p.word_wrap = True
    tf_p.paragraphs[0].text = "🖥️ Live Web Interface"
    tf_p.paragraphs[0].font.size = Pt(16)
    tf_p.paragraphs[0].font.bold = True
    tf_p.paragraphs[0].font.color.rgb = ACCENT_CYAN

    p = tf_p.add_paragraph()
    p.text = "\n• Dual Ingestion Modes:\n  1. Drag-and-drop .eml / .msg files\n  2. Paste raw RFC 822 email text\n\n• One-Click Demo Sample:\n  Instant live triage evaluation.\n\n• Live Deployment:\n  Deployed globally on Vercel Edge Serverless infrastructure."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 7: LIVE INCIDENT INVESTIGATION RESULTS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Demonstration: Automated Triage Results", category="DEMONSTRATION & PROTOTYPE")

    # Add investigation snapshot
    inv_img = "docs/images/investigation_results.png"
    if os.path.exists(inv_img):
        s7.shapes.add_picture(inv_img, Inches(0.8), Inches(1.7), Inches(7.5), Inches(5.0))

    # Findings Breakdown
    add_card(s7, Inches(8.6), Inches(1.7), Inches(3.9), Inches(5.0), border_color=ALERT_RED)
    tb_res = s7.shapes.add_textbox(Inches(8.8), Inches(1.9), Inches(3.5), Inches(4.6))
    tf_r = tb_res.text_frame
    tf_r.word_wrap = True
    tf_r.paragraphs[0].text = "⚡ Real Sample Output"
    tf_r.paragraphs[0].font.size = Pt(16)
    tf_r.paragraphs[0].font.bold = True
    tf_r.paragraphs[0].font.color.rgb = ALERT_RED

    p = tf_r.add_paragraph()
    p.text = "\n• Risk Verdict: 100/100 MALICIOUS\n• SPF / DKIM / DMARC: All Failed\n• Relay Origin: 185.220.101.45\n• Anchor Spoof: paypal.com text linking to 194.87.139.12\n• Payload: Invoice_OCT_2026.exe (PE Binary signature detected)\n• MITRE Techniques: 7 Mapped\n• Triage Time: < 1.2 seconds"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 8: MITRE ATT&CK & INCIDENT TIMELINE
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "MITRE ATT&CK Mapping & Chronological Timeline")

    # Left Card: MITRE
    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.0))
    tb_m = s8.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.2), Inches(4.6))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    tf_m.paragraphs[0].text = "🎯 MITRE ATT&CK® Enterprise Techniques"
    tf_m.paragraphs[0].font.size = Pt(15)
    tf_m.paragraphs[0].font.bold = True
    tf_m.paragraphs[0].font.color.rgb = PRIMARY_INDIGO

    p = tf_m.add_paragraph()
    p.text = "\n• T1566.001 - Spearphishing Attachment (Initial Access)\n• T1566.002 - Spearphishing Link (Initial Access)\n• T1036.005 - Masquerading Legitimate Name (Defense Evasion)\n• T1586.002 - Compromise Email Accounts (Resource Dev)\n• T1204.001 / .002 - User Execution of Link/File (Execution)\n• T1059 - Command & Scripting Interpreter (Execution)"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    # Right Card: Timeline
    add_card(s8, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb_t = s8.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.paragraphs[0].text = "⏱️ Chronological Chain of Custody"
    tf_t.paragraphs[0].font.size = Pt(15)
    tf_t.paragraphs[0].font.bold = True
    tf_t.paragraphs[0].font.color.rgb = ACCENT_CYAN

    p = tf_t.add_paragraph()
    p.text = "\n1. ORIGIN: Email dispatched from untrusted server.\n2. RELAY: Transit through intermediate mail relay hop.\n3. AUTH: SPF, DKIM, and DMARC evaluated.\n4. PAYLOAD: Hyperlink extracted and tested for spoofing.\n5. BINARY: Attachment decoded, hash and entropy computed.\n6. INTEL: Typosquatting / URLhaus threat cross-match.\n7. VERDICT: Formal incident status finalized."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 9: PROJECT PROGRESS & MILESTONES
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Project Progress, Milestones & Verification")

    milestones = [
        ("Phase 1: Foundation Engine", "RFC 822 parsing, hop tracing, SPF/DKIM/DMARC checks, anchor spoofing, and IOC defanging.", "100% COMPLETED", SUCCESS_GREEN),
        ("Phase 2: Threat Intelligence", "Levenshtein distance typosquatting engine, URLhaus malware feeds, AbuseIPDB integration.", "100% COMPLETED", SUCCESS_GREEN),
        ("Phase 3: Deep Inspection", "Magic bytes inspection (MZ/PDF/ZIP), Shannon entropy analysis, and MITRE ATT&CK mapping.", "100% COMPLETED", SUCCESS_GREEN),
        ("Phase 4: Case Management", "Case persistence (CASE-XXXXXXXX), analyst note logging, and standalone HTML audit report generator.", "100% COMPLETED", SUCCESS_GREEN),
        ("Production Deployment", "Full CI/CD integration, deployed globally on Vercel Serverless with GitHub tracking.", "100% COMPLETED", SUCCESS_GREEN)
    ]

    for i, (m_title, m_desc, m_stat, m_col) in enumerate(milestones):
        y = 1.7 + (i * 0.98)
        add_card(s9, Inches(0.8), Inches(y), Inches(11.7), Inches(0.85))
        tb = s9.shapes.add_textbox(Inches(1.0), Inches(y + 0.05), Inches(9.2), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = m_title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

        # Status badge on right
        sb = s9.shapes.add_textbox(Inches(10.3), Inches(y + 0.15), Inches(2.0), Inches(0.5))
        tf_s = sb.text_frame
        p_s = tf_s.paragraphs[0]
        p_s.text = m_stat
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = m_col

    # ==========================================
    # SLIDE 10: CONCLUSION & SUMMARY
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Summary & Live Demonstration Links", category="EVALUATION CONCLUSION")

    # Card 1: Key Achievements
    add_card(s10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.0))
    tb_c1 = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.2), Inches(4.6))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.paragraphs[0].text = "🏆 Key Achievements"
    tf_c1.paragraphs[0].font.size = Pt(16)
    tf_c1.paragraphs[0].font.bold = True
    tf_c1.paragraphs[0].font.color.rgb = ACCENT_CYAN

    p = tf_c1.add_paragraph()
    p.text = "\n• Practical SOC Problem Solved:\n  Automates manual triage, reducing analysis time from 15 mins to < 2 seconds.\n\n• Rigorous Security Standards:\n  Defangs all IOCs to prevent accidental execution.\n\n• Production-Ready Architecture:\n  FastAPI backend, responsive dark UI, automated pytest test suite (100% passing).\n\n• Formal Documentation:\n  Includes comprehensive Word specification (.docx) & live GitHub repository."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # Card 2: Links & Demo Information
    add_card(s10, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0), border_color=PRIMARY_INDIGO)
    tb_c2 = s10.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.paragraphs[0].text = "🔗 Project Links & Resources"
    tf_c2.paragraphs[0].font.size = Pt(16)
    tf_c2.paragraphs[0].font.bold = True
    tf_c2.paragraphs[0].font.color.rgb = PRIMARY_INDIGO

    p = tf_c2.add_paragraph()
    p.text = "\n• Live Deployed URL:\n  https://phishing-incident-investigation-wor.vercel.app\n\n• GitHub Repository:\n  https://github.com/Shivangi-gupta-loo1/\n  Phishing-Incident-Investigation-Workbench\n\n• Documentation Dossier:\n  Phishing_Investigation_Workbench_Documentation.docx\n\n• Ready for live demo questions & evaluation!"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE

    # Save presentation
    output_path = "/Users/shivika/Downloads/Project/MDC_Phishing_Workbench_Presentation.pptx"
    prs.save(output_path)
    print("Presentation successfully created at:", output_path)
    return output_path

if __name__ == "__main__":
    create_presentation()
